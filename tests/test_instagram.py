import sys
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook
import pytest

from influencer_service.deliverables import collect_workbook, write_results
from influencer_service.instagram import (
    InstagramAuthenticationError,
    InstaloaderClient,
    _login_error_message,
    _session_file_path,
    analyze_comments,
    collect_deliverable,
    shortcode_from_url,
)


class FakeClient:
    def fetch(self, shortcode):
        assert shortcode == "ABC_123-x"
        return {
            "likes": 120,
            "views": 900,
            "comments_count": 4,
            "comments": ["Amazing work!", "This is awful", "information"],
        }


def make_input(path, url="https://www.instagram.com/reel/ABC_123-x/"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Influencer Name", "Deliverable URL"])
    sheet.append(["Ada", url])
    workbook.save(path)


def test_shortcode_validation():
    assert shortcode_from_url("https://www.instagram.com/reel/ABC_123-x/?utm=x") == "ABC_123-x"
    with pytest.raises(ValueError):
        shortcode_from_url("https://example.com/reel/ABC/")
    with pytest.raises(ValueError):
        shortcode_from_url("https://www.instagram.com/profile/")


def test_sentiment_summary_ignores_blank_comments():
    summary = analyze_comments(["I love it", "terrible", "plain words", " "])
    assert summary.total_analyzed == 3
    assert summary.positive == 1
    assert summary.negative == 1
    assert summary.neutral == 1


def test_collect_one_deliverable():
    result = collect_deliverable("https://instagram.com/p/ABC_123-x/", FakeClient())
    assert result["likes"] == 120
    assert result["views"] == 900
    assert result["comments"] == 4
    assert result["comments_collected"] == 3


def test_workbook_batch_and_output(tmp_path):
    source = tmp_path / "input.xlsx"
    output = tmp_path / "output.xlsx"
    make_input(source)

    results = collect_workbook(source, FakeClient())
    write_results(output, results)

    assert results[0]["status"] == "ok"
    sheet = load_workbook(output).active
    assert sheet["A2"].value == "Ada"
    assert sheet["D2"].value == 120
    assert sheet["H2"].value == 1


def test_batch_records_bad_url_instead_of_aborting(tmp_path):
    source = tmp_path / "input.xlsx"
    make_input(source, "not-a-url")
    results = collect_workbook(source, FakeClient())
    assert results[0]["status"] == "error"
    assert "Instagram" in results[0]["error"]


class LoginException(Exception):
    pass


class InstaloaderException(Exception):
    pass


class FakeLoader:
    instances = []

    def __init__(self, **kwargs):
        self.loaded = None
        self.logged_in = None
        self.saved = None
        self.context = object()
        self.instances.append(self)

    def load_session_from_file(self, username, filename):
        self.loaded = (username, filename)

    def test_login(self):
        return "campaign_account"

    def login(self, username, password):
        self.logged_in = (username, password)

    def save_session_to_file(self, filename):
        self.saved = filename


def fake_instaloader_module():
    FakeLoader.instances.clear()
    return SimpleNamespace(
        Instaloader=FakeLoader,
        Post=SimpleNamespace(from_shortcode=lambda context, shortcode: None),
        exceptions=SimpleNamespace(
            LoginException=LoginException,
            InstaloaderException=InstaloaderException,
        ),
    )


def test_client_reuses_valid_authenticated_session(monkeypatch, tmp_path):
    module = fake_instaloader_module()
    session = tmp_path / "session"
    monkeypatch.setitem(sys.modules, "instaloader", module)
    monkeypatch.setenv("INSTAGRAM_USERNAME", "campaign_account")
    monkeypatch.setenv("INSTAGRAM_PASSWORD", "secret")
    monkeypatch.setenv("INSTAGRAM_SESSION_FILE", str(session))

    InstaloaderClient()

    loader = FakeLoader.instances[-1]
    assert loader.loaded == ("campaign_account", str(session))
    assert loader.logged_in is None


def test_client_logs_in_and_saves_when_session_is_missing(monkeypatch, tmp_path):
    module = fake_instaloader_module()
    session = tmp_path / "session"

    def missing_session(self, username, filename):
        raise FileNotFoundError(filename)

    monkeypatch.setattr(FakeLoader, "load_session_from_file", missing_session)
    monkeypatch.setitem(sys.modules, "instaloader", module)
    monkeypatch.setenv("INSTAGRAM_USERNAME", "campaign_account")
    monkeypatch.setenv("INSTAGRAM_PASSWORD", "secret")
    monkeypatch.setenv("INSTAGRAM_SESSION_FILE", str(session))

    InstaloaderClient()

    loader = FakeLoader.instances[-1]
    assert loader.logged_in == ("campaign_account", "secret")
    assert loader.saved == str(session)


def test_client_requires_login_credentials(monkeypatch):
    monkeypatch.setitem(sys.modules, "instaloader", fake_instaloader_module())
    monkeypatch.delenv("INSTAGRAM_USERNAME", raising=False)
    monkeypatch.delenv("INSTAGRAM_PASSWORD", raising=False)
    monkeypatch.delenv("INSTAGRAM_SESSION_FILE", raising=False)
    with pytest.raises(RuntimeError, match="login is required"):
        InstaloaderClient()


def test_checkpoint_error_has_clickable_url_and_retry_guidance():
    error = LoginException(
        "Login: Checkpoint required. Point your browser to "
        "/auth_platform/?apc=one-time-token - follow the instructions, then retry."
    )
    message = _login_error_message(error)
    assert "https://www.instagram.com/auth_platform/?apc=one-time-token" in message
    assert "run the collector again" in message
    assert "cannot be bypassed" in message
    assert "--browser-login" in message


def test_client_raises_specific_authentication_error(monkeypatch):
    module = fake_instaloader_module()

    def rejected_login(self, username, password):
        raise LoginException("Checkpoint required. /challenge/?token=abc")

    monkeypatch.setattr(FakeLoader, "login", rejected_login)
    monkeypatch.setitem(sys.modules, "instaloader", module)
    monkeypatch.setenv("INSTAGRAM_USERNAME", "campaign_account")
    monkeypatch.setenv("INSTAGRAM_PASSWORD", "secret")
    monkeypatch.delenv("INSTAGRAM_SESSION_FILE", raising=False)

    with pytest.raises(InstagramAuthenticationError, match="https://www.instagram.com/challenge"):
        InstaloaderClient()


def test_browser_login_mode_is_selected(monkeypatch, tmp_path):
    module = fake_instaloader_module()
    calls = []
    monkeypatch.setitem(sys.modules, "instaloader", module)
    monkeypatch.setenv("INSTAGRAM_USERNAME", "campaign_account")
    monkeypatch.setenv("INSTAGRAM_SESSION_FILE", str(tmp_path / "session"))
    monkeypatch.setattr(
        InstaloaderClient,
        "_authenticate_with_browser",
        lambda self, username, session: calls.append((username, session)),
    )

    InstaloaderClient(browser_login=True)

    assert calls == [("campaign_account", str(tmp_path / "session"))]


def test_session_directory_becomes_session_filename(tmp_path):
    assert _session_file_path("campaign_account", str(tmp_path)) == str(
        tmp_path / "session-campaign_account"
    )


def test_browser_cookie_transfer_sets_instaloader_username_before_save(monkeypatch, tmp_path):
    class Context:
        username = None

        def update_cookies(self, cookies):
            self.cookies = cookies

    class Loader:
        def __init__(self):
            self.context = Context()
            self.saved = None

        def test_login(self):
            return "campaign_account"

        def save_session_to_file(self, filename):
            if not self.context.username:
                raise LoginException("Login required")
            self.saved = filename

    class Browser:
        def get(self, url):
            self.url = url

        def get_cookies(self):
            return [{"name": "sessionid", "value": "approved-cookie"}]

        def quit(self):
            self.closed = True

    browser = Browser()
    webdriver = SimpleNamespace(
        ChromeOptions=lambda: SimpleNamespace(add_argument=lambda value: None),
        Chrome=lambda options: browser,
    )
    monkeypatch.setitem(sys.modules, "selenium", SimpleNamespace(webdriver=webdriver))
    monkeypatch.setattr("builtins.input", lambda prompt: "")
    client = object.__new__(InstaloaderClient)
    client._loader = Loader()

    session = str(tmp_path / "session")
    client._authenticate_with_browser("campaign_account", session)

    assert client._loader.context.username == "campaign_account"
    assert client._loader.saved == session
    assert browser.closed is True


def test_comment_endpoint_failure_preserves_engagement_metrics(monkeypatch):
    class Post:
        likes = 321
        video_view_count = 654
        is_video = True
        comments = 12

        def get_comments(self):
            raise InstaloaderException('200 OK - "fail" status')

    client = object.__new__(InstaloaderClient)
    client.max_comments = 500
    client._loader = SimpleNamespace(context=object())
    client._instaloader = SimpleNamespace(
        Post=SimpleNamespace(from_shortcode=lambda context, shortcode: Post()),
        exceptions=SimpleNamespace(InstaloaderException=InstaloaderException),
    )

    result = client.fetch("ABC")

    assert result["likes"] == 321
    assert result["views"] == 654
    assert result["comments_count"] == 12
    assert result["comments"] == []
    assert "sentiment is unavailable" in result["comments_error"]
